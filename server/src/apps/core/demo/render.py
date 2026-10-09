"""
Render authored Markdown course notes into real PDF / DOCX / TXT files.

Only the Markdown subset used by the demo library is supported: headings,
paragraphs, bullet and numbered lists, tables, fenced code, block quotes,
bold / italic / inline code.
"""

import html
import re
from pathlib import Path

import docx
import fitz
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

# ---------------------------------------------------------------------------
# Markdown → blocks
# ---------------------------------------------------------------------------


def parse_blocks(md: str) -> list[dict]:
    lines = md.replace("\r\n", "\n").split("\n")
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            body = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append({"type": "code", "text": "\n".join(body)})
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            blocks.append({"type": "h", "level": len(m.group(1)), "text": m.group(2).strip()})
            i += 1
            continue
        if re.match(r"^\s*\|.*\|\s*$", line) and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            def row(l):
                return [c.strip() for c in l.strip().strip("|").split("|")]
            head = row(line)
            i += 2
            rows = []
            while i < len(lines) and re.match(r"^\s*\|.*\|\s*$", lines[i]):
                rows.append(row(lines[i]))
                i += 1
            blocks.append({"type": "table", "head": head, "rows": rows})
            continue
        if re.match(r"^>\s?", line):
            body = []
            while i < len(lines) and re.match(r"^>\s?", lines[i]):
                body.append(re.sub(r"^>\s?", "", lines[i]))
                i += 1
            blocks.append({"type": "quote", "text": " ".join(body)})
            continue
        ul, ol = re.compile(r"^\s*[-*+]\s+"), re.compile(r"^\s*\d+[.)]\s+")
        if ul.match(line) or ol.match(line):
            ordered = bool(ol.match(line))
            rx = ol if ordered else ul
            items = []
            while i < len(lines) and (rx.match(lines[i]) or (items and re.match(r"^\s{2,}\S", lines[i]))):
                if rx.match(lines[i]):
                    items.append(rx.sub("", lines[i]))
                else:
                    items[-1] += " " + lines[i].strip()
                i += 1
            blocks.append({"type": "ol" if ordered else "ul", "items": items})
            continue
        para = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|```|>|\s*[-*+]\s|\s*\d+[.)]\s|\s*\|)", lines[i]):
            para.append(lines[i].strip())
            i += 1
        blocks.append({"type": "p", "text": " ".join(para)})
    return blocks


INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def inline_html(text: str) -> str:
    out = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            out.append(f"<code>{html.escape(part[1:-1])}</code>")
        elif part.startswith("**") and part.endswith("**"):
            out.append(f"<b>{html.escape(part[2:-2])}</b>")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append(f"<i>{html.escape(part[1:-1])}</i>")
        else:
            out.append(html.escape(part))
    return "".join(out)


def plain(text: str) -> str:
    return re.sub(r"\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`", lambda m: m.group(1) or m.group(2) or m.group(3), text)


# ---------------------------------------------------------------------------
# PDF (PyMuPDF Story — real pagination)
# ---------------------------------------------------------------------------

CSS = """
* { font-family: sans-serif; }
body { font-size: 10.5pt; line-height: 1.45; color: #1f1d1a; }
h1 { font-size: 19pt; margin: 0 0 4pt 0; color: #161412; }
h2 { font-size: 13.5pt; margin: 14pt 0 4pt 0; color: #161412; }
h3 { font-size: 11.5pt; margin: 10pt 0 3pt 0; color: #161412; }
p { margin: 0 0 7pt 0; }
ul, ol { margin: 0 0 7pt 0; }
li { margin: 0 0 2pt 0; }
code { font-family: monospace; font-size: 9.5pt; background-color: #f1eee8; }
pre { font-family: monospace; font-size: 9pt; background-color: #f4f2ee; padding: 6pt; margin: 0 0 8pt 0; }
table { border-collapse: collapse; margin: 2pt 0 9pt 0; }
th { background-color: #ece8e0; font-weight: bold; }
th, td { border: 0.5pt solid #b9b2a5; padding: 3pt 5pt; font-size: 9.5pt; text-align: left; }
blockquote { margin: 0 0 8pt 0; padding-left: 8pt; color: #4a463f; border-left: 2pt solid #8a6433; }
.meta { font-size: 9pt; color: #6b6559; margin: 0 0 12pt 0; }
"""


def to_html(blocks: list[dict], meta_line: str = "") -> str:
    parts = []
    for b in blocks:
        t = b["type"]
        if t == "h":
            lvl = min(3, b["level"])
            parts.append(f"<h{lvl}>{inline_html(b['text'])}</h{lvl}>")
            if lvl == 1 and meta_line:
                parts.append(f"<p class='meta'>{html.escape(meta_line)}</p>")
                meta_line = ""
        elif t == "p":
            parts.append(f"<p>{inline_html(b['text'])}</p>")
        elif t in ("ul", "ol"):
            items = "".join(f"<li>{inline_html(x)}</li>" for x in b["items"])
            parts.append(f"<{t}>{items}</{t}>")
        elif t == "code":
            parts.append(f"<pre>{html.escape(b['text'])}</pre>")
        elif t == "quote":
            parts.append(f"<blockquote>{inline_html(b['text'])}</blockquote>")
        elif t == "table":
            head = "".join(f"<th>{inline_html(c)}</th>" for c in b["head"])
            rows = "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in r) + "</tr>" for r in b["rows"])
            parts.append(f"<table><tr>{head}</tr>{rows}</table>")
    return "<body>" + "".join(parts) + "</body>"


def render_pdf(md: str, out: Path, meta_line: str = "", title: str = "") -> None:
    story = fitz.Story(html=to_html(parse_blocks(md), meta_line), user_css=CSS)
    writer = fitz.DocumentWriter(str(out))
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (56, 56, -56, -64)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()

    # Footer with page numbers + document metadata
    doc = fitz.open(str(out))
    total = doc.page_count
    for n, page in enumerate(doc, start=1):
        rect = page.rect
        page.insert_text((56, rect.height - 34), f"{meta_line}", fontsize=7.5, color=(0.45, 0.43, 0.39))
        label = f"Page {n} of {total}"
        page.insert_text((rect.width - 56 - fitz.get_text_length(label, fontsize=7.5), rect.height - 34), label, fontsize=7.5, color=(0.45, 0.43, 0.39))
    doc.set_metadata({"title": title, "author": meta_line.split(" · ")[0] if meta_line else "", "producer": "AuraStudy"})
    doc.saveIncr()
    doc.close()


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _runs(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def render_docx(md: str, out: Path, meta_line: str = "", title: str = "") -> None:
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    d.core_properties.title = title
    d.core_properties.author = meta_line.split(" · ")[0] if meta_line else ""
    first_heading = True
    for b in parse_blocks(md):
        t = b["type"]
        if t == "h":
            level = 0 if b["level"] == 1 else min(3, b["level"] - 1)
            d.add_heading(plain(b["text"]), level=level)
            if first_heading and meta_line:
                p = d.add_paragraph()
                r = p.add_run(meta_line)
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x6B, 0x65, 0x59)
            first_heading = False
        elif t == "p":
            _runs(d.add_paragraph(), b["text"])
        elif t in ("ul", "ol"):
            for item in b["items"]:
                _runs(d.add_paragraph(style="List Bullet" if t == "ul" else "List Number"), item)
        elif t == "code":
            p = d.add_paragraph()
            r = p.add_run(b["text"])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif t == "quote":
            p = d.add_paragraph(style="Intense Quote")
            _runs(p, b["text"])
        elif t == "table":
            table = d.add_table(rows=1, cols=len(b["head"]))
            table.style = "Light Grid Accent 1"
            for c, h in enumerate(b["head"]):
                table.rows[0].cells[c].text = plain(h)
            for r in b["rows"]:
                cells = table.add_row().cells
                for c, v in enumerate(r[: len(cells)]):
                    cells[c].text = plain(v)
            d.add_paragraph()
    for p in d.paragraphs:
        if p.style.name.startswith("Heading"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.save(str(out))


# ---------------------------------------------------------------------------
# TXT
# ---------------------------------------------------------------------------


def render_txt(md: str, out: Path, meta_line: str = "", title: str = "") -> None:
    lines = []
    for b in parse_blocks(md):
        t = b["type"]
        if t == "h":
            text = plain(b["text"])
            lines += ["", text.upper() if b["level"] == 1 else text, "=" * len(text) if b["level"] <= 2 else "", ""]
            if b["level"] == 1 and meta_line:
                lines += [meta_line, ""]
        elif t == "p":
            lines += [plain(b["text"]), ""]
        elif t in ("ul", "ol"):
            lines += [(f"{n}. " if t == "ol" else "- ") + plain(x) for n, x in enumerate(b["items"], 1)] + [""]
        elif t == "code":
            lines += ["    " + l for l in b["text"].split("\n")] + [""]
        elif t == "quote":
            lines += ["  > " + plain(b["text"]), ""]
        elif t == "table":
            lines += [" | ".join(plain(c) for c in b["head"])] + [" | ".join(plain(c) for c in r) for r in b["rows"]] + [""]
    out.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


RENDERERS = {"pdf": render_pdf, "docx": render_docx, "txt": render_txt}
